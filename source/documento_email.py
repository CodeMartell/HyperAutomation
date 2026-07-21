from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def gerar_ficha_cadastro(
    dados_cliente: dict,
    pasta_saida: str = "resources"
) -> Path:
    """
    Gera uma ficha de cadastro em formato Word.

    Args:
        dados_cliente: Dicionário contendo os dados extraídos do cliente.
        pasta_saida: Pasta em que o documento será salvo.

    Returns:
        Caminho completo do documento gerado.
    """

    campos_obrigatorios = [
        "nome",
        "sobrenome",
        "cpf",
        "email",
        "telefone",
        "data_nascimento",
        "endereco",
    ]

    campos_faltantes = [
        campo
        for campo in campos_obrigatorios
        if campo not in dados_cliente
    ]

    if campos_faltantes:
        raise ValueError(
            f"Campos obrigatórios ausentes: {', '.join(campos_faltantes)}"
        )

    raiz_projeto = Path(__file__).resolve().parent.parent
    diretorio_saida = raiz_projeto / pasta_saida
    diretorio_saida.mkdir(parents=True, exist_ok=True)

    data_hora = datetime.now().strftime("%Y%m%d_%H%M%S")
    caminho_documento = (
        diretorio_saida / f"ficha_cadastro_{data_hora}.docx"
    )

    documento = Document()

    estilo_normal = documento.styles["Normal"]
    estilo_normal.font.name = "Arial"
    estilo_normal.font.size = Pt(11)

    titulo = documento.add_heading(
        "FICHA DE CADASTRO DE CLIENTE",
        level=1
    )
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER

    empresa = documento.add_paragraph(
        "Portal Fake Soluções Digitais"
    )
    empresa.alignment = WD_ALIGN_PARAGRAPH.CENTER

    documento.add_paragraph(
        "Confira as informações abaixo e complete os campos necessários."
    )

    tabela = documento.add_table(rows=0, cols=2)
    tabela.style = "Table Grid"

    dados_formatados = [
        ("Nome", dados_cliente["nome"]),
        ("Sobrenome", dados_cliente["sobrenome"]),
        ("CPF", dados_cliente["cpf"]),
        ("E-mail", dados_cliente["email"]),
        ("Telefone", dados_cliente["telefone"]),
        ("Data de nascimento", dados_cliente["data_nascimento"]),
        ("Endereço", dados_cliente["endereco"]),
    ]

    for campo, valor in dados_formatados:
        linha = tabela.add_row().cells
        linha[0].text = campo
        linha[1].text = str(valor)

        linha[0].paragraphs[0].runs[0].bold = True

    documento.add_paragraph()

    documento.add_heading("Documentos necessários", level=2)

    documento.add_paragraph(
        "☐ Documento oficial com foto"
    )
    documento.add_paragraph(
        "☐ Comprovante de residência"
    )
    documento.add_paragraph(
        "☐ Ficha preenchida e convertida para PDF"
    )

    documento.add_paragraph()
    documento.add_paragraph(
        "Declaro que as informações apresentadas nesta ficha são verdadeiras."
    )

    documento.add_paragraph()
    documento.add_paragraph(
        "Assinatura: __________________________________________"
    )
    documento.add_paragraph(
        "Data: ______/______/________"
    )

    documento.save(caminho_documento)

    return caminho_documento


if __name__ == "__main__":
    dados_teste = {
        "nome": "Cliente",
        "sobrenome": "Teste",
        "cpf": "000.000.000-00",
        "email": "cliente.teste@email.com",
        "telefone": "(92) 99999-9999",
        "data_nascimento": "01/01/2000",
        "endereco": "Manaus - Amazonas",
    }

    try:
        arquivo_gerado = gerar_ficha_cadastro(dados_teste)
        print("Ficha gerada com sucesso.")
        print(f"Arquivo: {arquivo_gerado}")

    except Exception as erro:
        print(f"Erro ao gerar a ficha: {erro}")