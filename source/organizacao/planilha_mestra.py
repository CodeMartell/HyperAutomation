"""
Módulo de manipulação da Planilha Mestra — Processo 2 (Organização de Dados).

Lê, cria e atualiza a Planilha Mestra (Planilha_Mestra.xlsx), além de checar duplicidade.
"""

import logging
from pathlib import Path
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill

logger = logging.getLogger(__name__)

# Definição das colunas
COLUNAS = [
    "Data Processamento",
    "Nome",
    "Sobrenome",
    "CPF",
    "E-mail",
    "Telefone",
    "Data de Nascimento",
    "Endereço",
    "Nome Arquivo Ficha",
    "Hash Ficha",
    "Status Validação",
    "Erro Detalhado"
]

def inicializar_planilha(caminho_local: Path) -> None:
    """Cria a planilha com os cabeçalhos caso ela não exista."""
    if caminho_local.exists():
        return

    caminho_local.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fichas Processadas"

    # Estilizar cabeçalho
    ws.append(COLUNAS)
    font_cabecalho = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    fill_cabecalho = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")

    for col in range(1, len(COLUNAS) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = font_cabecalho
        cell.fill = fill_cabecalho

    wb.save(caminho_local)
    logger.info("Planilha Mestra inicializada localmente em: %s", caminho_local)

def checar_duplicado_hash(caminho_local: Path, hash_ficha: str) -> bool:
    """
    Verifica se o hash de uma ficha de cadastro já existe na planilha para evitar reprocessamento.
    """
    if not caminho_local.exists():
        return False

    try:
        wb = openpyxl.load_workbook(caminho_local, read_only=True)
        ws = wb["Fichas Processadas"]
        
        # Localiza o índice da coluna "Hash Ficha"
        hash_col_idx = COLUNAS.index("Hash Ficha") + 1

        for row in ws.iter_rows(min_row=2, values_only=True):
            if len(row) >= hash_col_idx:
                current_hash = row[hash_col_idx - 1]
                if current_hash == hash_ficha:
                    wb.close()
                    return True
        wb.close()
    except Exception as e:
        logger.error("Erro ao verificar duplicidade de hash na planilha: %s", e)
        
    return False

def adicionar_registro(
    caminho_local: Path,
    dados: dict,
    hash_ficha: str,
    status: str,
    erros: list[str],
    nome_ficha: str
) -> None:
    """
    Adiciona um registro (valido ou invalido) na Planilha Mestra.
    """
    inicializar_planilha(caminho_local)

    try:
        wb = openpyxl.load_workbook(caminho_local)
        ws = wb["Fichas Processadas"]

        data_processamento = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        msg_erros = "; ".join(erros) if erros else ""

        linha = [
            data_processamento,
            dados.get("nome", ""),
            dados.get("sobrenome", ""),
            dados.get("cpf", ""),
            dados.get("email", ""),
            dados.get("telefone", ""),
            dados.get("data_nascimento", ""),
            dados.get("endereco", ""),
            nome_ficha,
            hash_ficha,
            status,
            msg_erros
        ]

        ws.append(linha)
        
        # Estilizar linha de erro vs válida
        row_idx = ws.max_row
        if status == "ERRO":
            fill_error = PatternFill(start_color="F8CECC", end_color="F8CECC", fill_type="solid")
            for col in range(1, len(COLUNAS) + 1):
                ws.cell(row=row_idx, column=col).fill = fill_error

        wb.save(caminho_local)
        logger.info("Registro do cliente '%s %s' adicionado com status %s.", dados.get('nome'), dados.get('sobrenome'), status)
    except Exception as e:
        logger.error("Erro ao adicionar registro na planilha mestra: %s", e)
        raise e

# --- INTEGRAÇÃO GOOGLE DRIVE ---

def baixar_planilha_drive(servico, id_pasta_integrador: str, caminho_local: Path) -> str | None:
    """
    Tenta localizar e baixar a Planilha Mestra do Google Drive.
    Retorna o ID do arquivo no Drive, ou None se não existir.
    """
    import io
    from googleapiclient.http import MediaIoBaseDownload

    query = f"name = 'Planilha_Mestra.xlsx' and '{id_pasta_integrador}' in parents and trashed = false"
    try:
        resultados = servico.files().list(q=query, fields="files(id, name)").execute()
        arquivos = resultados.get("files", [])
        if not arquivos:
            return None

        file_id = arquivos[0]["id"]
        logger.info("[Drive] Planilha Mestra encontrada no Drive (ID: %s). Baixando...", file_id)

        request = servico.files().get_media(fileId=file_id)
        fh = io.BytesIO()
        downloader = MediaIoBaseDownload(fh, request)
        done = False
        while not done:
            status, done = downloader.next_chunk()

        caminho_local.parent.mkdir(parents=True, exist_ok=True)
        caminho_local.write_bytes(fh.getvalue())
        logger.info("[Drive] Planilha Mestra baixada com sucesso.")
        return file_id
    except Exception as e:
        logger.error("[Drive] Erro ao baixar planilha do Drive: %s", e)
        return None

def subir_planilha_drive(servico, id_pasta_integrador: str, caminho_local: Path, file_id: str | None = None) -> None:
    """
    Faz o upload da planilha atualizada de volta para o Google Drive.
    Se file_id for fornecido, atualiza o arquivo existente. Caso contrário, cria um novo.
    """
    from googleapiclient.http import MediaFileUpload

    media = MediaFileUpload(str(caminho_local), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", resumable=True)

    try:
        if file_id:
            logger.info("[Drive] Atualizando planilha Mestra no Drive (ID: %s)...", file_id)
            servico.files().update(
                fileId=file_id,
                media_body=media
            ).execute()
        else:
            logger.info("[Drive] Criando nova planilha Mestra no Drive...")
            metadados = {
                "name": "Planilha_Mestra.xlsx",
                "parents": [id_pasta_integrador]
            }
            servico.files().create(
                body=metadados,
                media_body=media,
                fields="id"
            ).execute()
        logger.info("[Drive] Sincronização da planilha mestra concluída.")
    except Exception as e:
        logger.error("[Drive] Erro ao fazer upload da planilha para o Drive: %s", e)
