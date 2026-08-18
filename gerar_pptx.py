"""
Gerador da Apresentação PowerPoint (Roteiro 13).
Cria apresentacao/Roteiro13_Apresentacao.pptx com os 18 slides.
"""

from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

RAIZ = Path(__file__).resolve().parent

def criar_apresentacao():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    slides_data = [
        ("SLIDE 1 — Capa", "Roteiro 13 – Integração, CI/CD, Containerização e Publicação no GHCR", "HyperAutomation Soluções Digitais | AX Academy / IFAM"),
        ("SLIDE 2 — Objetivo", "Objetivo da Solução Integrada", "Integrar os 5 processos em uma automação única, autônoma, testada e containerizada publicada no GHCR."),
        ("SLIDE 3 — Arquitetura", "Arquitetura Geral dos 5 Processos", "Processo 1 (Atendimento) -> Processo 2 (Organização) -> Processo 3 (Cadastro) -> Processo 4 (SAC) -> Processo 5 (Relatórios)"),
        ("SLIDE 4 — P1 & P2", "Processos 1 e 2 Reutilizados", "Reutilização dos módulos de leitura/validação de e-mails, extração PDF/OCR, Planilha Mestra Excel e arquivamento."),
        ("SLIDE 5 — P3 Cadastro", "Processo 3 — Setor de Cadastro", "Validação de dados, checagem prévia de duplicidades (CPF), integração com API externa e mecanismo de fallback."),
        ("SLIDE 6 — P4 SAC", "Processo 4 — Setor de SAC", "Protocolo SAC (SAC-YYYYMMDD-XXXX), notificação por e-mail/simulado, persistência JSON e medição de latência em ms."),
        ("SLIDE 7 — P5 Relatórios", "Processo 5 — Relatórios & Gerência", "Consolidação de KPIs globais, taxa % de sucesso, indicadores e relatórios gerenciais versionados."),
        ("SLIDE 8 — Integração", "Integração Real Entre os 5 Processos", "Passagem direta de payloads entre processos encadeados no ponto de entrada bot.py / pai.bot.py."),
        ("SLIDE 9 — Resiliência", "Logs, Erros, Fallback & Monitoramento", "Padrão de logs estruturados com timestamp, tratamento de exceções sem crash e fallbacks dedicados."),
        ("SLIDE 10 — Testes", "Testes Automatizados (pytest & pytest-cov)", "34 testes unitários e de integração aprovados (100% de sucesso) com 61% de cobertura global."),
        ("SLIDE 11 — Docker", "Docker & Containerização", "Dockerfile otimizado baseado em python:3.14-slim com dependências do Playwright e bot.py."),
        ("SLIDE 12 — CI/CD", "GitHub Actions (.github/workflows/ci-cd.yml)", "Esteira de CI (compileall + pytest) e CD (Docker build + GHCR push) condicional aos testes."),
        ("SLIDE 13 — GHCR", "Publicação no GitHub Container Registry", "Push automático da imagem ghcr.io/codemartell/hyperautomation:latest usando secrets.GITHUB_TOKEN."),
        ("SLIDE 14 — Validação GHCR", "Validação da Imagem Publicada", "docker pull ghcr.io/codemartell/hyperautomation:latest\ndocker run --rm ghcr.io/codemartell/hyperautomation:latest"),
        ("SLIDE 15 — Fluxo Completo", "Fluxo Completo da Atividade", "Desenvolvimento -> Testes -> Docker -> Git Push -> GitHub Actions -> Build GHCR -> Docker Pull -> Docker Run"),
        ("SLIDE 16 — Resultados", "Resultados Efetivamente Validados", "Todos os 5 processos integrados, testes 100% aprovados, Docker funcional e CI/CD configurado."),
        ("SLIDE 17 — Demonstração", "Roteiro Prático de Demonstração", "1. Executar bot.py\n2. Executar pytest -v\n3. Executar docker run\n4. Exibir GitHub Actions & GHCR"),
        ("SLIDE 18 — Conclusão", "Conclusão & Encerramento", "Solução 100% alinhada com todos os requisitos do Roteiro 13. Projeto concluído com sucesso!")
    ]

    # Cores
    COR_FUNDO = RGBColor(15, 23, 42)    # Dark slate
    COR_TEXTO = RGBColor(241, 245, 249) # Light
    COR_DESTAQUE = RGBColor(239, 35, 60) # Neon red

    for tag, titulo, corpo in slides_data:
        slide = prs.slides.add_slide(blank_layout)
        
        # Fundo escuro
        bg = slide.shapes.add_shape(1, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = COR_FUNDO
        bg.line.fill.background()

        # Tag superior
        tx_box_tag = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(0.5))
        tf_tag = tx_box_tag.text_frame
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = tag.upper()
        p_tag.font.size = Pt(14)
        p_tag.font.bold = True
        p_tag.font.color.rgb = COR_DESTAQUE
        p_tag.font.name = "Arial"

        # Título
        tx_box_t = slide.shapes.add_textbox(Inches(0.8), Inches(1.1), Inches(11.7), Inches(1.2))
        tf_t = tx_box_t.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = titulo
        p_t.font.size = Pt(32)
        p_t.font.bold = True
        p_t.font.color.rgb = COR_TEXTO
        p_t.font.name = "Arial"

        # Corpo
        tx_box_c = slide.shapes.add_textbox(Inches(0.8), Inches(2.5), Inches(11.7), Inches(4.2))
        tf_c = tx_box_c.text_frame
        tf_c.word_wrap = True
        p_c = tf_c.paragraphs[0]
        p_c.text = corpo
        p_c.font.size = Pt(20)
        p_c.font.color.rgb = RGBColor(203, 213, 225)
        p_c.font.name = "Arial"

    out_path = RAIZ / "apresentacao" / "Roteiro13_Apresentacao.pptx"
    prs.save(str(out_path))
    print(f"Apresentação PPTX gerada em: {out_path}")

if __name__ == "__main__":
    criar_apresentacao()
